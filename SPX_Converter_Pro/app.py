import io, os, re, zipfile, tempfile, pathlib, shutil, subprocess
from collections import Counter
from fastapi import FastAPI, UploadFile, File, Form, HTTPException
from fastapi.responses import FileResponse, StreamingResponse
from fastapi.staticfiles import StaticFiles
import fitz, pytesseract
from PIL import Image
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

app=FastAPI(title='SPX Converter Pro')
app.mount('/static',StaticFiles(directory='static'),name='static')
PAT=re.compile(r'SPX\s*VN\s*[0-9A-Z]{10,20}',re.I)
MAX_SIZE=60*1024*1024
MAX_FILES=1500

def extract_codes(text):
    return [re.sub(r'\s+','',m.group()).upper() for m in PAT.finditer(text)]

def read_document(path):
    suffix=path.suffix.lower()
    if suffix=='.pdf':
        doc=fitz.open(str(path)); output=[]
        if len(doc)>150: raise ValueError('PDF vượt quá 150 trang')
        for page in doc:
            output+=extract_codes(page.get_text())
            if not output or len(page.get_text().strip())<25:
                pix=page.get_pixmap(matrix=fitz.Matrix(2,2),alpha=False)
                im=Image.open(io.BytesIO(pix.tobytes('png')))
                output+=extract_codes(pytesseract.image_to_string(im,config='--psm 6'))
        return output
    if suffix in ('.jpg','.jpeg','.png','.webp'):
        with Image.open(path) as im:
            im.thumbnail((3500,3500))
            return extract_codes(pytesseract.image_to_string(im.convert('RGB'),config='--psm 11'))
    return []

def safe_target(root,name):
    parts=pathlib.PurePosixPath(name.replace('\\','/')).parts
    if not parts or any(x in ('..','.') for x in parts) or name.startswith(('/','\\')) or ':' in parts[0]:
        raise ValueError('Đường dẫn không hợp lệ trong file nén')
    target=root.joinpath(*parts)
    if not target.resolve().is_relative_to(root.resolve()): raise ValueError('Đường dẫn không hợp lệ')
    return target

def unpack(path,dest):
    suffix=path.suffix.lower()
    if suffix=='.zip':
        with zipfile.ZipFile(path) as z:
            infos=z.infolist()
            if len(infos)>MAX_FILES or sum(i.file_size for i in infos)>250*1024*1024: raise ValueError('File nén quá lớn')
            for info in infos:
                if info.is_dir(): continue
                target=safe_target(dest,info.filename)
                if target.suffix.lower() not in ('.pdf','.jpg','.jpeg','.png','.webp'): continue
                target.parent.mkdir(parents=True,exist_ok=True)
                with z.open(info) as src,open(target,'wb') as dst: shutil.copyfileobj(src,dst)
    elif suffix=='.rar':
        if not shutil.which('unar'): raise ValueError('Máy chủ chưa cài unar để giải nén RAR')
        proc=subprocess.run(['unar','-quiet','-force-overwrite','-output-directory',str(dest),str(path)],capture_output=True,text=True,timeout=90)
        if proc.returncode: raise ValueError('Không giải nén được RAR (kiểm tra mật khẩu hoặc file lỗi)')
        files=[p for p in dest.rglob('*') if p.is_file()]
        if len(files)>MAX_FILES or sum(p.stat().st_size for p in files)>250*1024*1024: raise ValueError('File nén quá lớn')
    else: raise ValueError('Định dạng không hỗ trợ')

@app.get('/')
def home(): return FileResponse('static/index.html')

@app.post('/convert')
async def convert(files:list[UploadFile]=File(...),deduplicate:bool=Form(True)):
    if len(files)>30: raise HTTPException(400,'Tối đa 30 file tải lên mỗi lần')
    rows=[]; errors=[]; seen=set(); total=0
    with tempfile.TemporaryDirectory() as temp:
        root=pathlib.Path(temp)
        for idx,f in enumerate(files):
            name=pathlib.Path(f.filename or f'file{idx}').name
            suffix=pathlib.Path(name).suffix.lower()
            if suffix not in ('.rar','.zip','.pdf','.jpg','.jpeg','.png','.webp'):
                errors.append(f'{name}: định dạng không hỗ trợ');continue
            data=await f.read(MAX_SIZE+1)
            if len(data)>MAX_SIZE: errors.append(f'{name}: vượt quá 60MB');continue
            source=root/f'upload_{idx}{suffix}';source.write_bytes(data)
            folder=root/f'extract_{idx}';folder.mkdir()
            try:
                if suffix in ('.rar','.zip'):
                    unpack(source,folder)
                    docs=[p for p in folder.rglob('*') if p.is_file() and p.suffix.lower() in ('.pdf','.jpg','.jpeg','.png','.webp')]
                else: docs=[source]
                for p in docs:
                    try:
                        codes=read_document(p)
                        total+=len(codes)
                        display=f'{name} / {p.relative_to(folder)}' if p!=source else name
                        for code in codes:
                            duplicate=code in seen
                            if not duplicate or not deduplicate: rows.append((code,display,'Trùng' if duplicate else 'Hợp lệ'))
                            seen.add(code)
                    except Exception as e: errors.append(f'{p.name}: {str(e)[:130]}')
            except Exception as e: errors.append(f'{name}: {str(e)[:130]}')
    wb=Workbook();ws=wb.active;ws.title='Danh sách mã SPX'
    ws.append(['STT','Mã vận đơn','File nguồn','Trạng thái'])
    for i,(code,source,status) in enumerate(rows,1):ws.append([i,code,source,status])
    ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
    for c in ws[1]:c.fill=PatternFill('solid',fgColor='153B68');c.font=Font(color='FFFFFF',bold=True);c.alignment=Alignment(horizontal='center')
    for col,width in {'A':9,'B':28,'C':65,'D':16}.items():ws.column_dimensions[col].width=width
    for row in ws.iter_rows(min_row=2):
        row[1].number_format='@'
        if row[3].value=='Trùng':
            for c in row:c.fill=PatternFill('solid',fgColor='FFF0D8')
    summary=wb.create_sheet('Thống kê');summary.append(['Chỉ số','Giá trị'])
    for k,v in [('Tổng lượt mã đọc được',total),('Mã duy nhất',len(seen)),('Mã trùng',total-len(seen)),('Số dòng xuất Excel',len(rows)),('Số lỗi/cảnh báo',len(errors))]:summary.append([k,v])
    summary.column_dimensions['A'].width=34;summary.column_dimensions['B'].width=20
    warn=wb.create_sheet('Cảnh báo');warn.append(['Thông báo'])
    for e in errors:warn.append([e])
    warn.column_dimensions['A'].width=110
    out=io.BytesIO();wb.save(out);out.seek(0)
    from urllib.parse import quote
    return StreamingResponse(out,media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',headers={'Content-Disposition':"attachment; filename=SPX_Converted.xlsx",'X-Total-Codes':str(total),'X-Unique-Codes':str(len(seen)),'X-Errors':str(len(errors)),'Access-Control-Expose-Headers':'X-Total-Codes, X-Unique-Codes, X-Errors'})
